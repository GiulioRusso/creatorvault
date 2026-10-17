"""Reading Meta Business Suite insights exports (CSV/XLSX)."""

from __future__ import annotations

import csv
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from creatorvault.models import log
from creatorvault.sources.instagram import fix_mojibake

# Possible column names, lowercased. Meta's export locale follows the
# account's UI language, not a fixed one, so each maps several observed
# header spellings.
COLUMN_ALIASES = {
    "publish_time": [
        "publish time", "ora di pubblicazione", "data di pubblicazione",
        "orario di pubblicazione", "date", "data",
    ],
    "permalink": ["permalink", "link permanente", "link", "url"],
    "description": ["description", "descrizione", "didascalia", "caption", "titolo del post"],
    "reach": ["reach", "copertura", "account raggiunti", "accounts reached"],
    "impressions": ["impressions", "impressioni", "visualizzazioni", "views"],
    "likes": ["likes", "mi piace", "like"],
    "comments": ["comments", "commenti"],
    "shares": ["shares", "condivisioni"],
    "saves": ["saves", "salvataggi", "elementi salvati"],
    "follows": ["follows", "nuovi follower", "segui"],
    "engagement": ["engagement", "interazioni", "interazioni con i contenuti"],
}


def load_table(path: Path) -> tuple[list[str], list[list[str]]]:
    """Loads CSV or XLSX, returning (header, rows)."""
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        # Meta exports in UTF-8, sometimes with a BOM
        with open(path, "r", encoding="utf-8-sig", newline="") as f:
            sample = f.read(8192)
            f.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            rows = [r for r in csv.reader(f, dialect) if any(c.strip() for c in r)]
    elif suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise SystemExit("[!] Needs openpyxl to read .xlsx files:  pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = []
        for r in ws.iter_rows(values_only=True):
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            rows.append(["" if c is None else str(c) for c in r])
    else:
        raise SystemExit(f"[!] Unsupported insights format: {suffix}")

    if not rows:
        raise SystemExit("[!] The insights file is empty.")

    # Meta sometimes prepends decorative header rows: look for the first row
    # that contains a column recognizable as a date.
    header_idx = 0
    for i, row in enumerate(rows[:10]):
        lowered = [c.strip().lower() for c in row]
        if any(alias in lowered for alias in COLUMN_ALIASES["publish_time"]):
            header_idx = i
            break
    return rows[header_idx], rows[header_idx + 1:]


def map_columns(header: list[str]) -> dict[str, int]:
    lowered = [h.strip().lower() for h in header]
    mapping: dict[str, int] = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias in lowered:
                mapping[canonical] = lowered.index(alias)
                break
        else:
            # partial match as a safety net
            for i, h in enumerate(lowered):
                if any(alias in h for alias in aliases):
                    mapping[canonical] = i
                    break
    return mapping


DATE_FORMATS = [
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d",
    "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y",
    "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %H:%M", "%m/%d/%Y",
    "%Y-%m-%dT%H:%M:%S",
]


def parse_datetime(value: str, tz: Any) -> datetime | None:
    value = value.strip()
    if not value:
        return None
    value = re.sub(r"\s*(UTC|GMT)[+-]?\d*$", "", value).strip()
    for fmt in DATE_FORMATS:
        try:
            naive = datetime.strptime(value, fmt)
            return naive.replace(tzinfo=tz).astimezone(timezone.utc)
        except ValueError:
            continue
    return None


def to_float(value: str) -> float:
    value = (value or "").strip().replace(".", "").replace(",", ".")
    value = re.sub(r"[^\d.\-]", "", value)
    try:
        return float(value)
    except ValueError:
        return 0.0


INSIGHTS_EXT = {".csv", ".tsv", ".txt", ".xlsx", ".xlsm"}


def discover_insight_files(insights_dir: Path) -> list[Path]:
    """One file per year inside insights_dir, auto-discovered."""
    if not insights_dir.exists():
        raise SystemExit(f"[!] Insights folder not found: {insights_dir}")
    files = sorted(p for p in insights_dir.iterdir() if p.suffix.lower() in INSIGHTS_EXT)
    if not files:
        raise SystemExit(
            f"[!] No insights file (.csv/.tsv/.xlsx) in {insights_dir}."
        )
    return files


def parse_insights(paths: list[Path], tz: Any) -> list[dict[str, Any]]:
    """
    Reads one or more insights files and merges their rows.
    Meta caps exports to short time windows (often a year), so it's normal
    to pass several. Duplicate rows across overlapping exports are dropped.
    """
    records: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()

    for path in paths:
        header, rows = load_table(path)
        cols = map_columns(header)
        if "publish_time" not in cols:
            raise SystemExit(
                f"[!] Can't find a publish-date column in {path.name}.\n"
                f"    Columns read: {header}\n"
                "    Rename the date column to 'Publish time' and try again."
            )

        added = 0
        for row in rows:
            def cell(name: str) -> str:
                i = cols.get(name)
                return row[i] if i is not None and i < len(row) else ""

            dt = parse_datetime(cell("publish_time"), tz)
            if dt is None:
                continue

            permalink = cell("permalink").strip() or None
            dedup_key = (permalink or "", dt.isoformat())
            if dedup_key in seen:
                continue
            seen.add(dedup_key)

            metrics = {
                m: to_float(cell(m))
                for m in ("reach", "impressions", "likes", "comments",
                          "shares", "saves", "follows", "engagement")
                if m in cols
            }
            if not metrics.get("engagement"):
                metrics["engagement"] = sum(
                    metrics.get(k, 0.0) for k in ("likes", "comments", "shares", "saves")
                )
            # Derived metric: interactions relative to reach. Lets posts from
            # different years be compared without penalizing ones published
            # when the page was smaller.
            if metrics.get("reach"):
                metrics["engagement_rate"] = round(
                    metrics["engagement"] / metrics["reach"] * 100, 2
                )

            records.append({
                "publish_time": dt,
                "permalink": permalink,
                "description": fix_mojibake(cell("description")),
                "metrics": metrics,
            })
            added += 1
        log(f"[i] {path.name}: {added} usable row(s).")

    log(f"[i] Total insights after dedup: {len(records)} row(s).")
    return records
